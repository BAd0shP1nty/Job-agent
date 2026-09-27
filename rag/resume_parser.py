"""Resume ingestion: safe upload validation, text extraction and local profile extraction.

Profile extraction is deterministic and local-first: every extracted skill is a
term that literally occurs in the resume, stored with the snippet that proves it.
Nothing is inferred or invented; the user reviews and corrects the profile in the
GUI before it is used.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from config.settings import ALLOWED_RESUME_EXTENSIONS, MAX_RESUME_BYTES
from database.models import CandidateProfile, ExperienceEntry
from screening.skill_filter import REQUIREMENT_VOCABULARY, _pattern


class ResumeValidationError(ValueError):
    pass


@dataclass
class ResumeSection:
    name: str
    text: str
    start: int
    end: int
    page: int | None = None


@dataclass
class ParsedResume:
    filename: str
    text: str
    sections: list[ResumeSection] = field(default_factory=list)
    headings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Upload validation
# ---------------------------------------------------------------------------


def validate_upload(filename: str, data: bytes) -> str:
    """Validate an uploaded resume; returns the normalized extension."""
    name = Path(filename or "").name
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_RESUME_EXTENSIONS:
        raise ResumeValidationError(
            f"Unsupported file type '{ext or 'none'}'. Upload a PDF, DOCX, TXT or MD resume."
        )
    if not data:
        raise ResumeValidationError("The uploaded file is empty.")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeValidationError(f"File is larger than {MAX_RESUME_BYTES // (1024 * 1024)} MB.")
    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise ResumeValidationError("File does not look like a valid PDF.")
    if ext == ".docx":
        if not data.startswith(b"PK"):
            raise ResumeValidationError("File does not look like a valid DOCX document.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
                if "word/document.xml" not in names:
                    raise ResumeValidationError("DOCX is missing its main document part.")
                if any(n.lower().endswith("vbaproject.bin") for n in names):
                    raise ResumeValidationError("Macro-enabled documents are not accepted.")
                total = sum(i.file_size for i in zf.infolist())
                if total > 50 * MAX_RESUME_BYTES:
                    raise ResumeValidationError("DOCX expands to an unreasonable size.")
        except zipfile.BadZipFile as exc:
            raise ResumeValidationError("DOCX archive is corrupt.") from exc
    if ext in {".txt", ".md"}:
        try:
            data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ResumeValidationError("Text resumes must be UTF-8 encoded.") from exc
    return ext


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------

_HEADING_WORDS = re.compile(
    r"^(professional )?(summary|profile|objective|experience|professional experience|work experience|employment"
    r"( history)?|education|certifications?|certifications? (&|and) professional development|skills|technical skills|"
    r"core competencies|technology( & development)? stack|projects|selected (ai )?projects|ai portfolio.*|"
    r"career focus|achievements|awards|languages|publications|interests|key skills|skills summary|areas of expertise|"
    r"technical expertise|core skills|tools|work history|career history|internships?|training|personal details|"
    r"declaration|references|hobbies|strengths|profile summary|executive summary|about me|ai / llm)$",
    re.IGNORECASE,
)


def _is_heading(line: str, style: str | None = None) -> bool:
    stripped = line.strip()
    if not stripped or len(stripped) > 60:
        return False
    if style and style.lower().startswith("heading 1"):
        return True
    return bool(_HEADING_WORDS.match(stripped.rstrip(":")))


def _extract_docx(data: bytes) -> tuple[list[str], list[str | None]]:
    try:
        import docx  # python-docx
    except ImportError:
        return _extract_docx_stdlib(data)

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    styles: list[str | None] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style is not None else None
        is_list = bool(style and style.lower().startswith("list")) or para._p.pPr is not None and \
            para._p.pPr.numPr is not None
        for piece in text.split("\n"):
            if piece.strip():
                lines.append(("• " if is_list and not piece.strip().startswith("•") else "") + piece.strip())
                styles.append(style)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
                styles.append(None)
    return lines, styles


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _extract_docx_stdlib(data: bytes) -> tuple[list[str], list[str | None]]:
    """Fallback DOCX reader using only the standard library (paragraphs, styles, bullets)."""
    import xml.etree.ElementTree as ET

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    lines: list[str] = []
    styles: list[str | None] = []
    for para in root.iter(f"{_W}p"):
        text = "".join(t.text or "" for t in para.iter(f"{_W}t")).strip()
        if not text:
            continue
        ppr = para.find(f"{_W}pPr")
        style_el = ppr.find(f"{_W}pStyle") if ppr is not None else None
        raw = style_el.get(f"{_W}val") if style_el is not None else None
        style = re.sub(r"(?<=[a-z])(?=[A-Z0-9])", " ", raw) if raw else None  # "Heading1" -> "Heading 1"
        is_list = (style or "").lower().startswith("list") or (ppr is not None and ppr.find(f"{_W}numPr") is not None)
        lines.append(("• " if is_list else "") + text)
        styles.append(style)
    return lines, styles


class ResumeLibraryError(ResumeValidationError):
    """A document-reading library could not be loaded on this computer."""


def _pdf_page_texts(data: bytes) -> list[str]:
    """Text of each PDF page. PyMuPDF first; pure-Python pypdf if PyMuPDF can't load
    (e.g. a missing Visual C++ runtime on Windows)."""
    errors = []
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz  # older PyMuPDF releases
        with fitz.open(stream=data, filetype="pdf") as doc:
            return [page.get_text("text") for page in doc]
    except ImportError as exc:
        errors.append(f"PyMuPDF: {exc}")
    try:
        from pypdf import PdfReader

        return [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]
    except ImportError as exc:
        errors.append(f"pypdf: {exc}")
    raise ResumeLibraryError(
        "This computer could not load a PDF reader (" + "; ".join(errors) + "). Run `pip install pypdf` in the "
        "app's virtual environment, or upload the resume as DOCX or TXT.")


def _extract_pdf(data: bytes) -> tuple[list[str], list[int]]:
    lines: list[str] = []
    pages: list[int] = []
    for page_number, page_text in enumerate(_pdf_page_texts(data), start=1):
        for line in page_text.splitlines():
            if line.strip():
                lines.append(line.strip())
                pages.append(page_number)
    # Remove running headers/footers: page numbers and lines repeated on several pages.
    page_count = len(set(pages))
    seen_on: dict[str, set[int]] = {}
    for line, page in zip(lines, pages):
        seen_on.setdefault(line, set()).add(page)
    keep = [
        i for i, line in enumerate(lines)
        if not re.fullmatch(r"(page\s*)?\d+(\s*(of|/)\s*\d+)?", line, re.IGNORECASE)
        and not (page_count >= 2 and len(seen_on[line]) >= 2 and len(line) < 120)
    ]
    return [lines[i] for i in keep], [pages[i] for i in keep]


def parse_resume(filename: str, data: bytes) -> ParsedResume:
    ext = validate_upload(filename, data)
    styles: list[str | None] = []
    pages: list[int | None] = []
    if ext == ".docx":
        lines, styles = _extract_docx(data)
        pages = [None] * len(lines)
    elif ext == ".pdf":
        lines, page_numbers = _extract_pdf(data)
        pages = list(page_numbers)
        styles = [None] * len(lines)
    else:
        raw_lines = [ln.strip() for ln in data.decode("utf-8").splitlines() if ln.strip()]
        styles = ["Heading 1" if ln.startswith("#") else None for ln in raw_lines]
        lines = [ln.lstrip("#").strip() for ln in raw_lines]
        pages = [None] * len(lines)
    if not lines:
        raise ResumeValidationError("No readable text was found. Scanned/image-only resumes are not supported yet.")

    text_parts: list[str] = []
    sections: list[ResumeSection] = []
    current_name, current_start, current_page = "Header", 0, pages[0] if pages else None
    offset = 0
    headings: list[str] = []
    for line, style, page in zip(lines, styles, pages):
        if _is_heading(line, style) and offset > 0:
            sections.append(ResumeSection(current_name, "\n".join(text_parts)[current_start:offset].strip(),
                                          current_start, offset, current_page))
            current_name, current_start, current_page = line.strip().rstrip(":").title(), offset, page
            headings.append(current_name)
        text_parts.append(line)
        offset = len("\n".join(text_parts)) + 1
    full_text = "\n".join(text_parts)
    sections.append(ResumeSection(current_name, full_text[current_start:].strip(), current_start, len(full_text),
                                  current_page))
    return ParsedResume(filename=Path(filename).name, text=full_text, sections=sections, headings=headings)


# ---------------------------------------------------------------------------
# Profile extraction (deterministic, evidence-backed)
# ---------------------------------------------------------------------------

_DATE_RANGE = re.compile(
    r"((jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{4}|\d{4})\s*[-–—to]+\s*"
    r"(present|current|now|(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{4}|\d{4})",
    re.IGNORECASE,
)
_YEARS_EXP = re.compile(r"(\d{1,2}\+?)\s*(\+\s*)?years?(\s+of)?\s+(professional\s+)?experience", re.IGNORECASE)
_KNOWN_LOCATIONS = ["Bangalore", "Bengaluru", "Hyderabad", "Pune", "Mumbai", "Delhi", "Gurgaon", "Gurugram",
                    "Noida", "Chennai", "Kolkata", "London", "Dublin", "Berlin", "Amsterdam", "Paris"]

EXTRA_RESUME_VOCABULARY = [
    "ITIL", "ITSM", "ServiceNow", "Dynatrace", "Elasticsearch", "AWS CloudWatch", "Amazon Bedrock", "SageMaker",
    "Microsoft Copilot Studio", "LangChain", "ReAct", "RAG", "MCP", "LLMs", "Prompt Engineering", "Agentic AI",
    "Generative AI", "Agentic Workflows", "Workflow Automation", "Python", "APIs", "Git", "Terraform", "Ansible",
    "CI/CD", "n8n", "Vercel", "Replit", "JupyterLab", "Jupyter Notebook", "Incident Management",
    "Change Management", "Change Governance", "Vendor Management", "Release Management", "Service Operations",
    "Service Delivery", "Service Management", "Stakeholder Management", "Root Cause Analysis", "RCA",
    "Online Charging", "OSS/BSS", "Mediation", "VoIP", "PSTN", "Broadband", "Telecom", "Agile", "SAFe", "Six Sigma",
    "SAP Ariba", "Procurement", "Project Management", "Program Management", "Machine Learning", "SRE", "DevOps",
    "Production Support", "Level-3 Support", "Regression Testing", "Test Cases", "Evaluation Harnesses",
    "Amazon Web Services", "AWS",
]


def _snippet(text: str, start: int, end: int, width: int = 90) -> str:
    s = max(0, text.rfind("\n", 0, start) + 1)
    e = text.find("\n", end)
    e = len(text) if e < 0 else e
    snippet = text[s:e].strip()
    if len(snippet) > 2 * width:
        rel = start - s
        snippet = snippet[max(0, rel - width): rel + width].strip()
    return snippet


def extract_skills(text: str) -> tuple[list[str], dict[str, str]]:
    """Skills that literally occur in the resume, with a proving snippet for each.

    Only literal occurrences count (no synonyms), and a term that only appears
    inside a longer extracted term (e.g. "Copilot Studio" inside "Microsoft
    Copilot Studio") is dropped to avoid double counting.
    """
    vocabulary = list(dict.fromkeys(EXTRA_RESUME_VOCABULARY + REQUIREMENT_VOCABULARY))
    found = [t for t in vocabulary if _pattern(t).search(text)]
    skills: list[str] = []
    evidence: dict[str, str] = {}
    for term in found:
        longer = [o for o in found if o != term and len(o) > len(term) and _pattern(term).search(o)]
        remainder = text
        for other in longer:
            remainder = _pattern(other).sub(" ", remainder)
        m = _pattern(term).search(remainder)
        if not m:
            continue
        original = _pattern(term).search(text)
        skills.append(term)
        evidence[term] = _snippet(text, original.start(), original.end())
    return skills, evidence


def _section(parsed: ParsedResume, *keywords: str) -> list[ResumeSection]:
    return [s for s in parsed.sections if any(k in s.name.lower() for k in keywords)]


BULLET_CHARS = ("•", "·", "◦", "‣", "▪", "■", "●", "\uf0b7", "\x7f", "-", "*", "–")


def _is_bullet(line: str) -> bool:
    return line.lstrip().startswith(BULLET_CHARS)


def _extract_experience(parsed: ParsedResume) -> list[ExperienceEntry]:
    """Find roles anchored on their date-range lines.

    Supported layouts (header lines directly above the dates line):
      "EMPLOYER | Title"            or      "EMPLOYER"
      "Month YYYY - Present | City"         "Title"
                                            "Month YYYY - Present | City"
    """
    entries: list[ExperienceEntry] = []
    for section in _section(parsed, "experience", "employment", "work history", "career history"):
        lines = [ln.strip() for ln in section.text.split("\n")[1:] if ln.strip()]
        anchors = [i for i, ln in enumerate(lines)
                   if not _is_bullet(ln) and len(ln) < 90 and (m := _DATE_RANGE.search(ln)) and m.start() < 3]
        blocks = []  # (header_start, date_idx, employer, title)
        for k, d in enumerate(anchors):
            floor = anchors[k - 1] + 1 if k else 0
            header = []
            j = d - 1
            while j >= floor and len(header) < 2 and not _is_bullet(lines[j]) and not _DATE_RANGE.search(lines[j]):
                header.insert(0, j)
                j -= 1
            if header and "|" in lines[header[-1]]:
                employer, _, title = (p.strip() for p in lines[header[-1]].partition("|"))
                header = [header[-1]]
            elif len(header) == 2:
                employer, title = lines[header[0]], lines[header[1]]
            elif header:
                employer, title = "", lines[header[0]]
            else:
                employer, title = "", ""
            blocks.append((header[0] if header else d, d, employer, title))
        for k, (start, d, employer, title) in enumerate(blocks):
            end = blocks[k + 1][0] if k + 1 < len(blocks) else len(lines)
            date_part, _, loc = (p.strip() for p in lines[d].partition("|"))
            highlights: list[str] = []
            for ln in lines[d + 1:end]:
                if _is_bullet(ln) or not highlights:
                    highlights.append(ln.lstrip("".join(BULLET_CHARS) + " ").strip())
                else:  # wrapped continuation line (common in PDFs)
                    highlights[-1] = f"{highlights[-1]} {ln}"
            if employer.isupper() and len(employer) > 4:
                employer = employer.title()
            entries.append(ExperienceEntry(employer=employer, title=title, dates=date_part, location=loc,
                                           highlights=highlights))
    return entries


def _bullets(section_list: list[ResumeSection]) -> list[str]:
    items: list[str] = []
    for section in section_list:
        for line in section.text.split("\n")[1:]:
            line = line.strip(" •-*\t")
            if line:
                items.append(line)
    return items


def extract_profile(parsed: ParsedResume) -> CandidateProfile:
    text = parsed.text
    lines = [ln for ln in text.split("\n") if ln.strip()]
    name = None
    if lines:
        first = lines[0].strip()
        if 1 < len(first.split()) <= 5 and not re.search(r"[@\d|:]", first):
            name = first.title() if first.isupper() else first
    headline = lines[1].strip() if len(lines) > 1 and "@" not in lines[1] else None

    header = parsed.sections[0].text if parsed.sections else text[:500]
    header_lines = [ln.strip() for ln in header.split("\n") if ln.strip()]
    location = None
    for i, ln in enumerate(header_lines[:-1]):  # "Location" label with the value on the next line
        if re.fullmatch(r"(current\s+)?(location|address|based in)\s*:?", ln, re.IGNORECASE):
            location = header_lines[i + 1]
            break
    if not location:
        m = re.search(r"(?:location|based in)\s*:\s*([^|\n]+)", header, re.IGNORECASE)
        location = m.group(1).strip() if m else None
    if not location:
        found = next((loc for loc in _KNOWN_LOCATIONS if re.search(rf"\b{loc}\b", header)), None)
        if found:
            m = re.search(rf"\b{found}\b[^|\n]*", header)
            location = m.group(0).strip() if m else found
    if not location:
        m = re.search(r"\b([A-Z][\w.'-]+(?: [A-Z][\w.'-]+)*, (?:[A-Z][\w.'-]+ ?)+, India)\b", header)
        location = m.group(1) if m else None

    years = _YEARS_EXP.search(text)
    skills, evidence = extract_skills(text)
    certifications = _bullets(_section(parsed, "certif"))
    education = _bullets(_section(parsed, "education"))
    return CandidateProfile(
        name=name,
        headline=headline,
        current_location=location,
        total_experience=(years.group(1) + " years") if years else None,
        skills=skills,
        skill_evidence=evidence,
        experience=_extract_experience(parsed),
        education=education,
        certifications=certifications,
        location_preferences=[location] if location else [],
        resume_sections=[{"name": s.name, "start": s.start, "end": s.end, "page": s.page} for s in parsed.sections],
    )
