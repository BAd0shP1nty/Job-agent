"""Shared pytest fixtures. All job data used in tests is mock data (see tests/fixtures)."""
from __future__ import annotations

import io
import shutil
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from config.settings import reset_settings  # noqa: E402
from database.connection import Database, set_database  # noqa: E402
from database.migrations import migrate  # noqa: E402


@pytest.fixture(autouse=True)
def _isolated_settings(tmp_path, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    reset_settings(anthropic_api_key=None, db_path=tmp_path / "test.db", data_dir=tmp_path)
    yield
    reset_settings()


@pytest.fixture
def db():
    database = Database(":memory:")
    migrate(database)
    set_database(database)
    yield database
    set_database(None)


@pytest.fixture
def file_db(tmp_path):
    path = tmp_path / "jobs.db"
    database = Database(path)
    migrate(database)
    yield database, path
    database.close()


@pytest.fixture
def skills_file(tmp_path):
    target = tmp_path / "list.py"
    shutil.copy(ROOT / "list.py", target)
    return target


def make_docx(paragraphs: list[tuple[str, str | None]]) -> bytes:
    import docx

    document = docx.Document()
    for text, style in paragraphs:
        if style:
            document.add_paragraph(text, style=style)
        else:
            document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


SAMPLE_RESUME = [
    ("JORDAN SAMPLE", None),  # fictional candidate
    ("Service Operations Lead | ITIL | Generative AI", None),
    ("Bangalore, India | jordan@example.com", None),
    ("PROFESSIONAL SUMMARY", "Heading 1"),
    ("Service operations leader with 12+ years of experience in incident management and ITSM using ServiceNow.",
     None),
    ("PROFESSIONAL EXPERIENCE", "Heading 1"),
    ("EXAMPLE TELECOM LTD | Service Delivery Manager", "Heading 2"),
    ("January 2018 - Present | Bengaluru", None),
    ("Own incident management and change governance for mobile services.", "List Bullet"),
    ("Built RAG prototypes with Python and LLMs for ticket triage.", "List Bullet"),
    ("EDUCATION", "Heading 1"),
    ("Bachelor of Engineering, Electronics | 2004 - 2008", None),
    ("CERTIFICATIONS", "Heading 1"),
    ("ITIL V3 Foundation", "List Bullet"),
]


@pytest.fixture
def sample_resume_bytes() -> bytes:
    return make_docx(SAMPLE_RESUME)

